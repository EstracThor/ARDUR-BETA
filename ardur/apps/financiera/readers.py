from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
import openpyxl
import xlrd
from django.core.exceptions import ValidationError
from .normalization import HEADERS

def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value

def formatted_identifier(cell):
    # Excel guarda algunos identificadores numéricos con una máscara de ceros.
    value = cell.value
    fmt = cell.number_format or ''
    if isinstance(value, (int, float)) and fmt and set(fmt) == {'0'}:
        return str(int(value)).zfill(len(fmt))
    return value

@contextmanager
def filas_excel(path, hoja):
    suffix = Path(path).suffix.lower()
    if suffix == '.xlsx':
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            if hoja not in book.sheetnames:
                raise ValidationError(f'Hoja inexistente. Disponibles: {", ".join(book.sheetnames)}')
            sheet = book[hoja]
            cells = sheet.iter_rows()
            header = [str(c.value or '').strip().upper() for c in next(cells)]
            validar_columnas(header)
            def rows():
                for number, row in enumerate(cells, 2):
                    values = [json_value(c.value) for c in row]
                    if all(v in (None, '') for v in values):
                        continue
                    raw = dict(zip(header, values))
                    normalized_source = dict(raw)
                    for key in ('CIU', 'CEDULA_RUC', 'CLAVE'):
                        normalized_source[key] = formatted_identifier(row[header.index(key)])
                    yield number, raw, normalized_source
            yield rows()
        finally:
            book.close()
    elif suffix == '.xls':
        book = xlrd.open_workbook(path, on_demand=True, formatting_info=True)
        try:
            if hoja not in book.sheet_names():
                raise ValidationError(f'Hoja inexistente. Disponibles: {", ".join(book.sheet_names())}')
            sheet = book.sheet_by_name(hoja)
            header = [str(v).strip().upper() for v in sheet.row_values(0)]
            validar_columnas(header)
            def rows():
                for index in range(1, sheet.nrows):
                    values = sheet.row_values(index)
                    if all(v in (None, '') for v in values):
                        continue
                    raw = dict(zip(header, values))
                    source = dict(raw)
                    for key in ('CIU', 'CEDULA_RUC', 'CLAVE'):
                        cell = sheet.cell(index, header.index(key))
                        fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                        if cell.ctype == xlrd.XL_CELL_NUMBER and fmt and set(fmt) == {'0'}:
                            source[key] = str(int(cell.value)).zfill(len(fmt))
                    yield index + 1, raw, source
            yield rows()
        finally:
            book.release_resources()
    else:
        raise ValidationError('Solo se admite .xls o .xlsx.')

def validar_columnas(header):
    missing = set(HEADERS) - set(header)
    if missing:
        raise ValidationError(f'Columnas faltantes: {", ".join(sorted(missing))}')
    if len(header) != len(set(header)):
        raise ValidationError('Hay encabezados duplicados; seleccione la hoja original de cartera.')
